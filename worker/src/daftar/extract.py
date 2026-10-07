"""Turn a sniffed file into raw rows. Amounts stay strings until a profile is confirmed."""

from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass, field
from datetime import datetime

import fitz
from PIL import Image, ImageOps

from daftar.decode import decode_bytes
from daftar.errors import CodedError
from daftar.mapping import detect_date_format, detect_separators, field_for_header, propose_column_map
from daftar.money import MoneyError, decimal_string_from_excel, parse_amount
from daftar.preview import render_table_png, union_bbox
from daftar.textclean import text_layer_unreliable


@dataclass
class Token:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float
    confidence: float | None = None


@dataclass
class RawRow:
    row_no: int
    source_page: int | None
    source_bbox: dict | None
    cells: dict[str, str]
    confidence: float | None = None


@dataclass
class PageImage:
    page_no: int
    png: bytes
    width: int
    height: int


@dataclass
class Extraction:
    headers: list[str]
    rows: list[RawRow]
    opening_minor: int | None = None
    closing_minor: int | None = None
    period_start: str | None = None
    period_end: str | None = None
    currency: str | None = None
    page_count: int = 1
    methods: list[tuple[int, str]] = field(default_factory=list)
    pages: list[PageImage] = field(default_factory=list)
    encoding: str | None = None
    decimal_sep: str = "."
    thousands_sep: str = ","
    date_format: str = "%Y-%m-%d"
    kind: str = "csv"
    # Per-page text used by field extraction (text layer and/or OCR). Never drop OCR here.
    raw_pages: list[str] = field(default_factory=list)


OPENING_RE = re.compile(r"opening(?:\s+balance)?\s*[:|,]?\s*([-(]?\d[\d.,]*\)?)", re.I)
CLOSING_RE = re.compile(r"closing(?:\s+balance)?\s*[:|,]?\s*([-(]?\d[\d.,]*\)?)", re.I)
PERIOD_RE = re.compile(
    r"period\s*[:|,]?\s*(\d{4}-\d{2}-\d{2}|\d{2}[/-]\d{2}[/-]\d{4})\s*(?:to|-)\s*(\d{4}-\d{2}-\d{2}|\d{2}[/-]\d{2}[/-]\d{4})",
    re.I,
)
CURRENCY_RE = re.compile(r"currency\s*[:|,]\s*([A-Za-z]{3})", re.I)


def extract_file(content: bytes, kind: str, password: str | None = None) -> Extraction:
    if kind == "csv":
        return _extract_csv(content)
    if kind == "xlsx":
        return _extract_xlsx(content)
    if kind == "xls":
        return _extract_xls(content)
    if kind == "pdf":
        return _extract_pdf(content, password)
    if kind == "image":
        return _extract_image(content)
    raise CodedError("unsupported_type")


def _extract_csv(content: bytes) -> Extraction:
    text, encoding = decode_bytes(content)
    lines = text.splitlines()
    header_at = _header_index(lines)
    if header_at is None:
        raise CodedError("header_not_found")
    delimiter = _delimiter(lines[header_at : header_at + 6])
    meta = _metadata_from_lines(lines[:header_at], delimiter)
    import csv

    reader = csv.reader(io.StringIO("\n".join(lines[header_at:])), delimiter=delimiter)
    table = [row for row in reader if any(cell.strip() for cell in row)]
    if not table:
        raise CodedError("header_not_found")
    headers = [cell.strip() for cell in table[0]]
    rows = []
    for offset, cols in enumerate(table[1:], start=1):
        cells = {headers[i]: cols[i].strip() if i < len(cols) else "" for i in range(len(headers)) if headers[i]}
        rows.append(RawRow(offset, 1, None, cells, 1.0))
    extraction = _finish_tabular(headers, rows, meta, encoding, "csv")
    extraction.methods = [(1, "csv")]
    return extraction


def _extract_xlsx(content: bytes) -> Extraction:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheet = _pick_sheet(workbook)
    grid = [[_excel_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return _from_grid(grid, "xlsx")


def _extract_xls(content: bytes) -> Extraction:
    import pandas as pd

    frame = pd.read_excel(io.BytesIO(content), header=None, engine="xlrd")
    grid = [[_excel_cell(value) for value in row] for row in frame.itertuples(index=False, name=None)]
    return _from_grid(grid, "xlsx")


def _from_grid(grid: list[list[str]], kind: str) -> Extraction:
    lines = [",".join(cell for cell in row) for row in grid]
    header_at = None
    for index, row in enumerate(grid[:40]):
        if _row_is_header(row):
            header_at = index
            break
    if header_at is None:
        raise CodedError("header_not_found")
    meta_lines = []
    for row in grid[:header_at]:
        text = " ".join(cell for cell in row if cell).strip()
        if text:
            meta_lines.append(text if ":" in text or "," in text else text)
    meta = _metadata_from_free_text("\n".join(meta_lines))
    # Also accept two-column key/value title rows.
    for row in grid[:header_at]:
        if len(row) >= 2 and row[0] and row[1]:
            _assign_meta(meta, row[0], row[1])
    headers = [cell.strip() for cell in grid[header_at]]
    rows = []
    number = 0
    for raw in grid[header_at + 1 :]:
        if not any(cell.strip() for cell in raw):
            continue
        number += 1
        cells = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            cells[header] = raw[index].strip() if index < len(raw) else ""
        rows.append(RawRow(number, 1, None, cells, 1.0))
    extraction = _finish_tabular(headers, rows, meta, None, kind)
    extraction.methods = [(1, "xlsx")]
    return extraction


def _finish_tabular(headers, rows, meta, encoding, kind) -> Extraction:
    amount_samples = _amount_samples(headers, rows)
    date_samples = _date_samples(headers, rows)
    decimal_sep, thousands_sep = detect_separators(amount_samples)
    date_format = detect_date_format(date_samples)
    opening = _money(meta.get("opening"), decimal_sep, thousands_sep)
    closing = _money(meta.get("closing"), decimal_sep, thousands_sep)
    png, boxes, width, height = render_table_png(headers, [{"row_no": row.row_no, "cells": row.cells} for row in rows])
    by_row = {box["row_no"]: box["bbox"] for box in boxes}
    for row in rows:
        row.source_bbox = by_row.get(row.row_no)
        row.source_page = 1
    return Extraction(
        headers=[header for header in headers if header],
        rows=rows,
        opening_minor=opening,
        closing_minor=closing,
        period_start=meta.get("period_start"),
        period_end=meta.get("period_end"),
        currency=(meta.get("currency") or "AED").upper(),
        page_count=1,
        pages=[PageImage(1, png, width, height)],
        encoding=encoding,
        decimal_sep=decimal_sep,
        thousands_sep=thousands_sep,
        date_format=date_format,
        kind=kind,
    )


def _max_pdf_pages() -> int:
    try:
        return max(1, int(os.environ.get("BOOKS_EXTRACT_MAX_PAGES", "50")))
    except ValueError:
        return 50


def _extract_pdf(content: bytes, password: str | None) -> Extraction:
    document = fitz.open(stream=content, filetype="pdf")
    try:
        if document.needs_pass:
            if not password:
                raise CodedError("password_required")
            if not document.authenticate(password):
                raise CodedError("bad_password")
        max_pages = _max_pdf_pages()
        if document.page_count > max_pages:
            raise CodedError(f"too_many_pages:{document.page_count}>{max_pages}")
        extraction = Extraction(headers=[], rows=[], page_count=document.page_count, kind="pdf")
        meta_text = []
        next_row = 1
        for index, page in enumerate(document, start=1):
            text = page.get_text("text") or ""
            if _page_has_text(text):
                words = page.get_text("words") or []
                zoom = 1400 / max(page.rect.width, 1)
                matrix = fitz.Matrix(zoom, zoom)
                pixmap = page.get_pixmap(matrix=matrix, alpha=False)
                tokens = [
                    Token(word[4], word[0] * zoom, word[1] * zoom, word[2] * zoom, word[3] * zoom, 1.0)
                    for word in words
                ]
                headers, rows = _rows_from_tokens(tokens, index, pixmap.width, pixmap.height)
                for row in rows:
                    row.row_no = next_row
                    next_row += 1
                if not extraction.headers:
                    extraction.headers = headers
                extraction.rows.extend(rows)
                extraction.pages.append(PageImage(index, pixmap.tobytes("png"), pixmap.width, pixmap.height))
                extraction.methods.append((index, "pdf_text"))
                extraction.raw_pages.append(text)
                meta_text.append(text)
            else:
                zoom = 2
                pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                headers, rows, png, width, height, ocr_text = _ocr_image(image, index)
                for row in rows:
                    row.row_no = next_row
                    next_row += 1
                if not extraction.headers:
                    extraction.headers = headers
                extraction.rows.extend(rows)
                extraction.pages.append(PageImage(index, png, width, height))
                engine = getattr(recognize, "last_engine", "ocr")
                extraction.methods.append((index, engine))
                extraction.raw_pages.append(ocr_text)
                meta_text.append(ocr_text)
        meta = _metadata_from_free_text("\n".join(meta_text))
        _apply_meta(extraction, meta)
        return extraction
    finally:
        document.close()


def _extract_image(content: bytes) -> Extraction:
    image = Image.open(io.BytesIO(content))
    image = ImageOps.exif_transpose(image).convert("RGB")
    headers, rows, png, width, height, ocr_text = _ocr_image(image, 1)
    for number, row in enumerate(rows, start=1):
        row.row_no = number
    engine = getattr(recognize, "last_engine", "ocr")
    extraction = Extraction(
        headers=headers,
        rows=rows,
        page_count=1,
        pages=[PageImage(1, png, width, height)],
        methods=[(1, engine)],
        kind="image",
        raw_pages=[ocr_text],
    )
    _apply_meta(extraction, _metadata_from_free_text(ocr_text))
    return extraction


def _ocr_image(image: Image.Image, page_no: int):
    processed, tokens = recognize(image)
    width, height = processed.size
    buffer = io.BytesIO()
    processed.save(buffer, format="PNG")
    headers, rows = _rows_from_tokens(tokens, page_no, width, height)
    ocr_text = " ".join(token.text for token in tokens)
    return headers, rows, buffer.getvalue(), width, height, ocr_text


def recognize(image: Image.Image) -> tuple[Image.Image, list[Token]]:
    """OCR via docTR (primary) with EasyOCR Arabic/low-confidence fallback."""
    provider = os.environ.get("OCR_PROVIDER", "local")
    if provider == "textract":
        raise CodedError("textract_not_available")
    if provider != "local":
        raise CodedError("unknown_ocr_provider")
    from daftar.ocr_engine import recognize_image

    processed = _preprocess(image)
    result = recognize_image(processed)
    # Stash engine on module for callers that inspect last OCR route.
    recognize.last_engine = result.engine  # type: ignore[attr-defined]
    recognize.last_models = result.model_names  # type: ignore[attr-defined]
    tokens = [
        Token(t.text, t.x0, t.y0, t.x1, t.y1, t.confidence) for t in result.tokens
    ]
    return processed, tokens


recognize.last_engine = "none"  # type: ignore[attr-defined]
recognize.last_models = []  # type: ignore[attr-defined]


def _preprocess(image: Image.Image) -> Image.Image:
    gray = ImageOps.autocontrast(image.convert("L"))
    gray = _deskew(gray)
    if gray.width < 1400:
        scale = 1400 / gray.width
        gray = gray.resize((int(gray.width * scale), int(gray.height * scale)), Image.Resampling.LANCZOS)
    return gray.convert("RGB")


def _deskew(image: Image.Image) -> Image.Image:
    import numpy as np

    gray = np.asarray(image)
    best_angle = 0.0
    best_score = -1.0
    for angle in np.linspace(-6, 6, 25):
        rotated = image.rotate(float(angle), expand=False, fillcolor=255)
        ink = np.asarray(rotated) < 180
        score = float(ink.sum(axis=1).var())
        if score > best_score:
            best_score = score
            best_angle = float(angle)
    if abs(best_angle) < 0.4:
        return image
    return image.rotate(best_angle, expand=True, fillcolor=255)


def _rows_from_tokens(tokens: list[Token], page_no: int, page_width: int, page_height: int):
    lines = _cluster_lines(tokens)
    header_index = None
    ranges: list[tuple[str, float, float]] = []
    headers: list[str] = []
    for index, line in enumerate(lines):
        mapped = []
        seen: set[str] = set()
        for token in line:
            field = field_for_header(token.text)
            if field and field not in seen:
                mapped.append((token, field))
                seen.add(field)
        if "txn_date" in seen and len(seen) >= 3 and ({"debit", "credit", "balance"} & seen):
            header_index = index
            ordered = sorted(mapped, key=lambda item: item[0].x0)
            centers = [((token.x0 + token.x1) / 2, token.text.strip()) for token, _field in ordered]
            headers = [name for _center, name in centers]
            for position, (center, name) in enumerate(centers):
                left = 0 if position == 0 else (centers[position - 1][0] + center) / 2
                right = 1_000_000 if position == len(centers) - 1 else (center + centers[position + 1][0]) / 2
                ranges.append((name, left, right))
            break
    if header_index is None:
        return [], []
    rows: list[RawRow] = []
    for line in lines[header_index + 1 :]:
        cells = {name: "" for name in headers}
        used: list[Token] = []
        for token in line:
            center = (token.x0 + token.x1) / 2
            for name, left, right in ranges:
                if left <= center < right:
                    cells[name] = (cells[name] + " " + token.text).strip()
                    used.append(token)
                    break
        if not any(cells.values()):
            continue
        # A continuation line has no date-like token. Append it to the previous description.
        date_header = next((name for name in headers if field_for_header(name) == "txn_date"), None)
        description_header = next((name for name in headers if field_for_header(name) == "description"), None)
        date_text = cells.get(date_header or "", "")
        if rows and date_header and not re.search(r"\d", date_text) and description_header and cells.get(description_header):
            previous = rows[-1]
            previous.cells[description_header] = (
                previous.cells.get(description_header, "") + " " + cells[description_header]
            ).strip()
            continue
        confidence_values = [token.confidence for token in used if token.confidence is not None]
        confidence = sum(confidence_values) / len(confidence_values) if confidence_values else None
        rows.append(
            RawRow(
                row_no=len(rows) + 1,
                source_page=page_no,
                source_bbox=union_bbox(used, page_width, page_height),
                cells=cells,
                confidence=confidence,
            )
        )
    return headers, rows


def _cluster_lines(tokens: list[Token]) -> list[list[Token]]:
    if not tokens:
        return []
    heights = [max(1.0, token.y1 - token.y0) for token in tokens]
    tolerance = max(6.0, sorted(heights)[len(heights) // 2] * 0.6)
    ordered = sorted(tokens, key=lambda token: ((token.y0 + token.y1) / 2, token.x0))
    lines: list[list[Token]] = [[ordered[0]]]
    current_y = (ordered[0].y0 + ordered[0].y1) / 2
    for token in ordered[1:]:
        mid = (token.y0 + token.y1) / 2
        if abs(mid - current_y) <= tolerance:
            lines[-1].append(token)
            current_y = sum((item.y0 + item.y1) / 2 for item in lines[-1]) / len(lines[-1])
        else:
            lines.append([token])
            current_y = mid
    for line in lines:
        line.sort(key=lambda token: token.x0)
    return lines


def _page_has_text(text: str) -> bool:
    minimum = int(os.environ.get("PDF_PROBE_MIN_CHARS", "30"))
    if text_layer_unreliable(text):
        return False  # unmapped-font garbage -> OCR route
    compact = re.sub(r"\s+", "", text or "")
    if len(compact) < minimum:
        return False
    has_date = re.search(r"\d{4}-\d{2}-\d{2}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4}", text) is not None
    has_amount = re.search(r"\d+[.,]\d{2}", text) is not None
    return has_date and has_amount


def _apply_meta(extraction: Extraction, meta: dict[str, str]) -> None:
    samples = _amount_samples(extraction.headers, extraction.rows)
    date_samples = _date_samples(extraction.headers, extraction.rows)
    extraction.decimal_sep, extraction.thousands_sep = detect_separators(samples)
    extraction.date_format = detect_date_format(date_samples)
    extraction.opening_minor = _money(meta.get("opening"), extraction.decimal_sep, extraction.thousands_sep)
    extraction.closing_minor = _money(meta.get("closing"), extraction.decimal_sep, extraction.thousands_sep)
    extraction.period_start = meta.get("period_start")
    extraction.period_end = meta.get("period_end")
    extraction.currency = (meta.get("currency") or "AED").upper()
    if not extraction.headers:
        extraction.headers = list(propose_column_map([]).keys())


def _header_index(lines: list[str]) -> int | None:
    for index, line in enumerate(lines[:40]):
        if _row_is_header(_split_loose(line)):
            return index
    return None


def _split_loose(line: str) -> list[str]:
    if "\t" in line:
        return line.split("\t")
    if line.count(";") > line.count(","):
        return line.split(";")
    return line.split(",")


def _row_is_header(cells: list[str]) -> bool:
    fields = {field_for_header(cell) for cell in cells if cell and field_for_header(cell)}
    return "txn_date" in fields and len(fields) >= 3 and bool(fields & {"debit", "credit", "balance"})


def _delimiter(lines: list[str]) -> str:
    header = lines[0] if lines else ""
    counts = {",": header.count(","), "\t": header.count("\t"), ";": header.count(";")}
    best = max(counts, key=counts.get)
    return best if counts[best] else ","


def _metadata_from_lines(lines: list[str], delimiter: str) -> dict[str, str]:
    meta: dict[str, str] = {}
    for line in lines:
        if not line.strip():
            continue
        if delimiter in line:
            parts = [part.strip() for part in line.split(delimiter)]
            if len(parts) >= 2:
                _assign_meta(meta, parts[0], parts[1])
                if len(parts) >= 3 and parts[0].lower().startswith("period"):
                    meta["period_end"] = _normalise_period_token(parts[2])
            continue
        if ":" in line:
            key, value = line.split(":", 1)
            _assign_meta(meta, key, value)
    # Free-text fallback for "Period: A to B" held in one cell.
    extra = _metadata_from_free_text("\n".join(lines))
    for key, value in extra.items():
        meta.setdefault(key, value)
    return meta


def _metadata_from_free_text(text: str) -> dict[str, str]:
    meta: dict[str, str] = {}
    opening = OPENING_RE.search(text or "")
    closing = CLOSING_RE.search(text or "")
    period = PERIOD_RE.search(text or "")
    currency = CURRENCY_RE.search(text or "")
    if opening:
        meta["opening"] = opening.group(1)
    if closing:
        meta["closing"] = closing.group(1)
    if period:
        meta["period_start"] = _normalise_period_token(period.group(1))
        meta["period_end"] = _normalise_period_token(period.group(2))
    if currency:
        meta["currency"] = currency.group(1).upper()
    return meta


def _assign_meta(meta: dict[str, str], key: str, value: str) -> None:
    name = re.sub(r"\s+", " ", key.strip().casefold())
    value = value.strip()
    if name in {"opening", "opening balance"}:
        meta["opening"] = value
    elif name in {"closing", "closing balance"}:
        meta["closing"] = value
    elif name in {"period start", "from"}:
        meta["period_start"] = _normalise_period_token(value)
    elif name in {"period end", "to"}:
        meta["period_end"] = _normalise_period_token(value)
    elif name == "period":
        match = re.search(r"(\d{4}-\d{2}-\d{2}).*?(\d{4}-\d{2}-\d{2})", value)
        if match:
            meta["period_start"], meta["period_end"] = match.group(1), match.group(2)
    elif name == "currency":
        meta["currency"] = value.upper()


def _normalise_period_token(value: str) -> str:
    text = value.strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text


def _amount_samples(headers: list[str], rows: list[RawRow]) -> list[str]:
    names = [header for header in headers if field_for_header(header) in {"debit", "credit", "balance"}]
    samples = []
    for row in rows:
        for name in names:
            if row.cells.get(name):
                samples.append(row.cells[name])
    return samples


def _date_samples(headers: list[str], rows: list[RawRow]) -> list[str]:
    names = [header for header in headers if field_for_header(header) == "txn_date"]
    return [row.cells.get(name, "") for row in rows for name in names if row.cells.get(name)]


def _money(raw: str | None, decimal_sep: str, thousands_sep: str) -> int | None:
    if raw is None:
        return None
    try:
        return parse_amount(raw, decimal_sep=decimal_sep, thousands_sep=thousands_sep)
    except (MoneyError, TypeError):
        return None


def _excel_cell(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    try:
        import pandas as pd

        if value is pd.NaT or (not isinstance(value, str) and pd.isna(value)):
            return ""
    except Exception:
        pass
    try:
        import numpy as np

        if isinstance(value, np.floating):
            value = float(value)
        elif isinstance(value, np.integer):
            value = int(value)
    except Exception:
        pass
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, float):
        if value != value:  # NaN
            return ""
        try:
            return decimal_string_from_excel(value)
        except MoneyError:
            return ""
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if text.lower() in {"nan", "nat"}:
        return ""
    return text


def _pick_sheet(workbook):
    for name in workbook.sheetnames:
        sheet = workbook[name]
        if sheet.sheet_state == "visible":
            return sheet
    return workbook[workbook.sheetnames[0]]
